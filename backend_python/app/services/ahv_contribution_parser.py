from __future__ import annotations

import io
import re
from dataclasses import dataclass
from datetime import date
from typing import Iterable

from openpyxl import load_workbook
import pdfplumber


@dataclass
class AhvContributionRow:
    year: int
    annual_income_chf: float | None
    contribution_total_chf: float | None
    age_end_of_year: int | None
    ahv_rent_share_chf: float | None
    source: str = "manual"


def _parse_float(value: object) -> float | None:
    if value is None or value == "":
        return None
    cleaned = str(value).strip().replace("'", "").replace(" ", "").replace(",", ".")
    if cleaned == "":
        return None
    try:
        parsed = float(cleaned)
    except ValueError:
        return None
    return parsed


def _parse_int(value: object) -> int | None:
    parsed = _parse_float(value)
    if parsed is None:
        return None
    return int(parsed)


def parse_ahv_contributions_from_xlsx(data: bytes) -> list[AhvContributionRow]:
    workbook = load_workbook(filename=io.BytesIO(data), data_only=True, read_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []

    header_row_index = _find_header_row_index(rows)
    if header_row_index is None:
        return []

    headers = [str(col or "").strip().lower() for col in rows[header_row_index]]
    idx_year = _find_header_index(headers, ("jahr", "year"))
    idx_income = _find_header_index(headers, ("jahreseinkommen", "income"))
    idx_total = _find_header_index(headers, ("total", "beitragssumme", "beitrag"))
    idx_age = _find_header_index(headers, ("alter", "age"))
    idx_share = _find_header_index(headers, ("ahv rente anteil", "rente anteil"))

    parsed: list[AhvContributionRow] = []
    for row in rows[header_row_index + 1 :]:
        year = _parse_int(row[idx_year]) if idx_year is not None and idx_year < len(row) else None
        if year is None or year < 1900:
            continue
        parsed.append(
            AhvContributionRow(
                year=year,
                annual_income_chf=_parse_float(row[idx_income]) if idx_income is not None and idx_income < len(row) else None,
                contribution_total_chf=_parse_float(row[idx_total]) if idx_total is not None and idx_total < len(row) else None,
                age_end_of_year=_parse_int(row[idx_age]) if idx_age is not None and idx_age < len(row) else None,
                ahv_rent_share_chf=_parse_float(row[idx_share]) if idx_share is not None and idx_share < len(row) else None,
                source="xlsx",
            )
        )
    return parsed


def parse_ahv_contributions_from_text(text: str) -> list[AhvContributionRow]:
    parsed: list[AhvContributionRow] = []
    for line in text.splitlines():
        normalized = _normalize_pdf_line(line)
        row = _parse_ahv_row_line(normalized)
        if row is not None:
            parsed.append(row)
            continue

        # Fallback for pipe or tab separated content.
        parts = [chunk for chunk in re.split(r"\s{2,}|\t|\|", normalized.strip()) if chunk]
        if len(parts) >= 3:
            year = _parse_int(parts[0])
            if year is not None and year >= 1900:
                parsed.append(
                    AhvContributionRow(
                        year=year,
                        annual_income_chf=_parse_float(parts[1]) if len(parts) > 1 else None,
                        contribution_total_chf=_parse_float(parts[2]) if len(parts) > 2 else None,
                        age_end_of_year=_parse_int(parts[3]) if len(parts) > 3 else None,
                        ahv_rent_share_chf=_parse_float(parts[4]) if len(parts) > 4 else None,
                        source="pdf",
                    )
                )
    return parsed


def parse_ahv_contributions_from_pdf(data: bytes) -> list[AhvContributionRow]:
    combined_text = _extract_pdf_text_and_tables(data)
    return parse_ahv_contributions_from_text(combined_text)


def parse_ahv_pdf_import(data: bytes) -> tuple[list[AhvContributionRow], dict[str, object]]:
    combined_text = _extract_pdf_text_and_tables(data)
    rows = parse_ahv_contributions_from_text(combined_text)
    metadata = _extract_ahv_import_metadata(combined_text, rows)
    return rows, metadata


def import_ahv_data_from_pdf(pdf_path: str) -> list[dict]:
    """
    Extract structured AHV contribution data from PDF.
    Returns list of dictionaries with contribution records.
    """
    with open(pdf_path, "rb") as handle:
        rows = parse_ahv_contributions_from_pdf(handle.read())

    return [
        {
            "year": row.year,
            "annual_income_chf": row.annual_income_chf,
            "contribution_total_chf": row.contribution_total_chf,
            "age_end_of_year": row.age_end_of_year,
            "ahv_rent_share_chf": row.ahv_rent_share_chf,
            "source": row.source,
        }
        for row in rows
    ]


def validate_ahv_contributions(
    rows: Iterable[AhvContributionRow],
    birth_date: date | None = None,
) -> list[dict[str, object]]:
    issues: list[dict[str, object]] = []
    for row in rows:
        if row.annual_income_chf is not None and row.annual_income_chf < 0:
            issues.append({"year": row.year, "code": "NEGATIVE_INCOME", "message": "Jahreseinkommen darf nicht negativ sein."})
        if row.contribution_total_chf is not None and row.contribution_total_chf < 0:
            issues.append({"year": row.year, "code": "NEGATIVE_CONTRIBUTION", "message": "Beitragssumme darf nicht negativ sein."})
        if birth_date is not None:
            min_year = birth_date.year + 16
            if row.year < min_year:
                issues.append({"year": row.year, "code": "YEAR_BEFORE_ELIGIBLE_AGE", "message": "Jahr liegt vor möglichem AHV-Beitragsbeginn."})

    sorted_rows = sorted(rows, key=lambda value: value.year)
    for prev, nxt in zip(sorted_rows, sorted_rows[1:]):
        if nxt.year > prev.year + 1:
            issues.append({"year": prev.year, "code": "CONTRIBUTION_GAP", "message": f"Mögliche Lücke zwischen {prev.year} und {nxt.year}."})
    return issues


def _find_header_index(headers: list[str], candidates: tuple[str, ...]) -> int | None:
    for idx, header in enumerate(headers):
        for candidate in candidates:
            if candidate in header:
                return idx
    return None


def _find_header_row_index(rows: list[tuple[object, ...]]) -> int | None:
    for idx, row in enumerate(rows[:20]):
        headers = [str(col or "").strip().lower() for col in row]
        if (
            _find_header_index(headers, ("jahr", "year")) is not None
            and _find_header_index(headers, ("jahreseinkommen", "income")) is not None
        ):
            return idx
    return None


def _normalize_pdf_line(line: str) -> str:
    normalized = line.replace("\xa0", " ").strip()
    if not normalized:
        return ""

    # Remove row index prefixes like "0 1968 ..." / "1 1994 ..."
    normalized = re.sub(r"^\d+\s+(?=\d{4}\b)", "", normalized)
    normalized = normalized.replace("CHF", " ")
    normalized = normalized.replace("%", " ")
    normalized = normalized.replace("’", "'")
    return re.sub(r"\s+", " ", normalized).strip()


def _parse_ahv_row_line(line: str) -> AhvContributionRow | None:
    if not line:
        return None
    # Skip headers and obvious non-data rows.
    lowered = line.lower()
    if "jahr" in lowered and "jahreseinkommen" in lowered:
        return None
    if "initial entry" in lowered:
        return None

    numeric_tokens = re.findall(r"-?[0-9][0-9'.,]*", line)
    if len(numeric_tokens) < 4:
        return None

    year_index: int | None = None
    for idx, token in enumerate(numeric_tokens):
        token_clean = token.replace("'", "")
        if "." in token_clean or "," in token_clean:
            continue
        try:
            year_candidate = int(token_clean)
        except ValueError:
            continue
        if 1900 <= year_candidate <= 2200:
            year_index = idx
            break

    if year_index is None:
        return None
    year = int(numeric_tokens[year_index].replace("'", ""))
    if year < 1900 or year > 2200:
        return None

    values = numeric_tokens[year_index + 1 :]
    if len(values) < 3:
        return None

    income = _parse_float(values[0])
    total = _parse_float(values[1])
    if income is None and total is None:
        return None

    age: int | None = None
    share: float | None = None

    # Variant A: <year> <income> <total> <age> <share> ...
    # Variant B: <year> <income> <total> <share> <monthly_rent> (age missing)
    if len(values) >= 4:
        third = _parse_float(values[2])
        fourth = _parse_float(values[3])
        third_is_int = third is not None and float(third).is_integer()
        third_int = int(third) if third_is_int else None
        if third_int is not None and 15 <= third_int <= 120 and fourth is not None:
            age = third_int
            share = fourth
        else:
            share = third
    else:
        share = _parse_float(values[2])

    return AhvContributionRow(
        year=year,
        annual_income_chf=income,
        contribution_total_chf=total,
        age_end_of_year=age,
        ahv_rent_share_chf=share,
        source="pdf",
    )


def _extract_pdf_text_and_tables(data: bytes) -> str:
    extracted_text_parts: list[str] = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text() or ""
            if page_text:
                extracted_text_parts.append(page_text)

            table = page.extract_table()
            if table:
                for cells in table:
                    row_text = " | ".join(str(cell or "").strip() for cell in cells)
                    if row_text.strip():
                        extracted_text_parts.append(row_text)
    return "\n".join(extracted_text_parts)


def _extract_ahv_import_metadata(
    source_text: str,
    rows: list[AhvContributionRow],
) -> dict[str, object]:
    lines = [line.strip() for line in source_text.splitlines() if line.strip()]
    income_values = [row.annual_income_chf for row in rows if row.annual_income_chf is not None]
    contribution_values = [row.contribution_total_chf for row in rows if row.contribution_total_chf is not None]

    average_income_reference_chf = _extract_reference_amount(
        source_text,
        (
            r"durchschnitt(?:seinkommen)?",
            r"durchschnittliches\s+einkommen",
        ),
    )
    total_accumulated_reference_chf = _extract_reference_amount(
        source_text,
        (
            r"total(?:\s+akkumuliert)?",
            r"summe(?:\s+beitr[a-z]*)?",
        ),
    )
    latest_monthly_rent_reference_chf = _extract_last_monthly_rent(lines)
    difference_table_preview = _extract_difference_table(lines)

    return {
        "rows_detected": len(rows),
        "computed_average_income_chf": (sum(income_values) / len(income_values)) if income_values else None,
        "computed_total_contributions_chf": sum(contribution_values) if contribution_values else None,
        "average_income_reference_chf": average_income_reference_chf,
        "total_accumulated_reference_chf": total_accumulated_reference_chf,
        "latest_monthly_rent_reference_chf": latest_monthly_rent_reference_chf,
        "difference_table": difference_table_preview,
        "detected_fields": [
            "year",
            "annual_income_chf",
            "contribution_total_chf",
            "age_end_of_year",
            "ahv_rent_share_chf",
        ],
    }


def _extract_reference_amount(source_text: str, keyword_patterns: tuple[str, ...]) -> float | None:
    for pattern in keyword_patterns:
        for match in re.finditer(
            rf"{pattern}[^\n\r0-9-]{{0,40}}(-?[0-9][0-9'., ]+)",
            source_text,
            flags=re.IGNORECASE,
        ):
            parsed = _parse_float(match.group(1))
            if parsed is not None:
                return parsed
    return None


def _extract_last_monthly_rent(lines: list[str]) -> float | None:
    last_value: float | None = None
    for line in lines:
        if "rente monat" not in line.lower() and "monat" not in line.lower():
            continue
        for token in re.findall(r"-?[0-9][0-9'.,]*", line):
            parsed = _parse_float(token)
            if parsed is not None:
                last_value = parsed
    return last_value


def _extract_difference_table(lines: list[str]) -> list[dict[str, object]]:
    entries: list[dict[str, object]] = []
    for line in lines:
        if "differenz" not in line.lower() and "delta" not in line.lower():
            continue
        values = [_parse_float(token) for token in re.findall(r"-?[0-9][0-9'.,]*", line)]
        numeric_values = [value for value in values if value is not None]
        if not numeric_values:
            continue
        entries.append(
            {
                "label": line[:120],
                "delta_monthly_chf": numeric_values[-1],
            }
        )
        if len(entries) >= 10:
            break
    return entries
